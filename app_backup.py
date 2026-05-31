from flask import Flask, render_template, request, redirect, session, jsonify
import sqlite3
import json
import os
import re
from datetime import datetime
from werkzeug.utils import secure_filename
import sys
import urllib.request

sys.path.append(os.path.join(os.path.dirname(os.path.abspath(__file__)), "utils"))

try:
    from ai_compare import compare_candidates
except ImportError:
    def compare_candidates(candidate_list):
        return {"error": "ai_compare not found", "ranking": [], "total_candidates": 0,
                "hire_count": 0, "shortlist_count": 0, "reject_count": 0}

app = Flask(__name__)
app.secret_key = "lumina_secret_key_2026"
app.config['SESSION_PERMANENT'] = True

UPLOAD_FOLDER = "uploads"
ALLOWED_EXTENSIONS = {"pdf", "docx", "txt"}
GEMINI_KEY = "AIzaSyABYO9FWgLlN-BvDB9MU8eBEbU88zryUpQ"

def get_db():
    conn = sqlite3.connect("users.db")
    conn.row_factory = sqlite3.Row
    return conn

def get_screening_db():
    conn = sqlite3.connect("screening_system.db")
    conn.row_factory = sqlite3.Row
    return conn

def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS

def extract_resume_data(filepath, filename):
    raw_text = ""
    try:
        if filename.endswith(".pdf"):
            import pdfplumber
            with pdfplumber.open(filepath) as pdf:
                for page in pdf.pages:
                    raw_text += page.extract_text() or ""
        elif filename.endswith(".docx"):
            from docx import Document
            doc = Document(filepath)
            for para in doc.paragraphs:
                raw_text += para.text + "\n"
        elif filename.endswith(".txt"):
            with open(filepath, "r", encoding="utf-8") as f:
                raw_text = f.read()
    except Exception as e:
        print("Resume read error:", e)
    return raw_text

def call_gemini(prompt):
    payload = json.dumps({
        "contents": [{"role": "user", "parts": [{"text": prompt}]}],
        "generationConfig": {"maxOutputTokens": 1000, "temperature": 0.7}
    }).encode("utf-8")
    url = "https://generativelanguage.googleapis.com/v1beta/models/gemini-2.5-flash:generateContent?key=" + GEMINI_KEY
    req = urllib.request.Request(url, data=payload, headers={"Content-Type": "application/json"}, method="POST")
    try:
        with urllib.request.urlopen(req) as r:
            result = json.loads(r.read().decode())
            return result["candidates"][0]["content"]["parts"][0]["text"]
    except urllib.error.HTTPError as e:
        error = e.read().decode()
        print("Gemini error:", error)
        return None
    except Exception as e:
        print("Gemini exception:", e)
        return None

@app.route("/test-ai")
def test_ai():
    result = call_gemini("Say hello in one sentence.")
    return result or "Gemini call failed"

@app.route("/")
def home():
    return render_template("public/landing_page_1.html")
# ============================================================
# REPLACE your existing login, signup, logout routes with these
# ============================================================



@app.route("/dashboard")
def dashboard():
    if "user" not in session:
        return redirect("/login")
    conn = get_screening_db()
    try:
        total = conn.execute("SELECT COUNT(*) FROM candidates").fetchone()[0]
        hired = conn.execute("SELECT COUNT(*) FROM candidates WHERE status='hired'").fetchone()[0]
        shortlisted = conn.execute("SELECT COUNT(*) FROM candidates WHERE status='shortlisted'").fetchone()[0]
        rejected = conn.execute("SELECT COUNT(*) FROM candidates WHERE status='rejected'").fetchone()[0]
        recent = [dict(r) for r in conn.execute("SELECT * FROM candidates ORDER BY created_at DESC LIMIT 5").fetchall()]
    except Exception as e:
        print("Dashboard error:", e)
        total = hired = shortlisted = rejected = 0
        recent = []
    conn.close()
    return render_template("dashboard_layer/recruiter-dashboard.html",
                           total=total, hired=hired, shortlisted=shortlisted, rejected=rejected, recent=recent)

@app.route("/upload", methods=["GET", "POST"])
def upload():
    if "user" not in session:
        return redirect("/login")
    if request.method == "POST":
        job_role = request.form.get("job_role")
        file = request.files.get("resume")
        if file and allowed_file(file.filename):
            filename = secure_filename(file.filename)
            os.makedirs(UPLOAD_FOLDER, exist_ok=True)
            file.save(os.path.join(UPLOAD_FOLDER, filename))
            session["uploaded_file"] = filename
            session["job_role"] = job_role
            return redirect("/screening")
        return render_template("resume_pipeline_layer/upload-resume-1.html", error="Invalid file!")
    return render_template("resume_pipeline_layer/upload-resume-1.html")

@app.route("/upload-2", methods=["GET", "POST"])
def upload2():
    if "user" not in session:
        return redirect("/login")
    if request.method == "POST":
        file = request.files.get("resume")
        if file and allowed_file(file.filename):
            filename = secure_filename(file.filename)
            os.makedirs(UPLOAD_FOLDER, exist_ok=True)
            file.save(os.path.join(UPLOAD_FOLDER, filename))
            session["uploaded_file"] = filename
            return redirect("/screening")
        return render_template("resume_pipeline_layer/upload-resume-2.html", error="Invalid file!")
    return render_template("resume_pipeline_layer/upload-resume-2.html")

@app.route("/screening")
def screening():
    if "user" not in session:
        return redirect("/login")
    filename = session.get("uploaded_file", "")
    job_role = session.get("job_role", "General Role")
    candidate_name = "No Resume Uploaded"
    candidate_email = "N/A"
    candidate_phone = "N/A"
    skills = []
    ats_score = 0

    if filename:
        filepath = os.path.join(UPLOAD_FOLDER, filename)
        raw_text = extract_resume_data(filepath, filename)
        if raw_text:
            email_match = re.search(r'[\w\.-]+@[\w\.-]+\.\w+', raw_text)
            if email_match:
                candidate_email = email_match.group()
            phone_match = re.search(r'[\+\d][\d\s\-\(\)]{9,15}', raw_text)
            if phone_match:
                candidate_phone = phone_match.group().strip()
            lines = [l.strip() for l in raw_text.split('\n') if l.strip()]
            if lines:
                candidate_name = lines[0]
            known_skills = ['python','javascript','react','node','sql','aws','docker','kubernetes',
                           'java','typescript','flask','django','mongodb','postgresql','git',
                           'machine learning','tensorflow','pytorch','agile','devops','linux',
                           'html','css','vue','angular','redis','tableau','power bi','excel','c++','c#','php','ruby']
            text_lower = raw_text.lower()
            skills_raw = [s for s in known_skills if s in text_lower]
            ats_score = min(95, 40 + len(skills_raw) * 5 + (10 if candidate_email != "N/A" else 0))
            conn = get_screening_db()
            try:
                existing = conn.execute("SELECT id FROM candidates WHERE email=?", (candidate_email,)).fetchone()
                if not existing:
                    conn.execute("""INSERT INTO candidates
                        (name, email, phone, skills_json, score, status, notes, created_at, updated_at, user_id)
                        VALUES (?,?,?,?,?,?,?,?,?,?)""",
                        (candidate_name, candidate_email, candidate_phone, json.dumps(skills_raw),
                         ats_score, 'pending', f'Auto-screened for {job_role}',
                         datetime.now().isoformat(), datetime.now().isoformat(), session.get("user", "unknown")))
                    conn.commit()
                    print("Candidate saved to DB!")
                else:
                    conn.execute("""UPDATE candidates SET name=?, phone=?, skills_json=?, score=?, updated_at=? WHERE email=?""",
                                 (candidate_name, candidate_phone, json.dumps(skills_raw), ats_score, datetime.now().isoformat(), candidate_email))
                    conn.commit()
            except Exception as e:
                print("DB error:", e)
            conn.close()
            skills = [{"name": s.title(), "score": min(98, 70 + i * 4)} for i, s in enumerate(skills_raw[:8])]

    decision = "HIRE" if ats_score >= 75 else "SHORTLIST" if ats_score >= 55 else "REJECT"
    insight = f"{candidate_name} screened for {job_role}. Found {len(skills)} matching skills. ATS: {ats_score}%. Decision: {decision}."
    return render_template("resume_pipeline_layer/ai_screening.html",
                           candidate_name=candidate_name, candidate_email=candidate_email,
                           candidate_phone=candidate_phone, skills=skills,
                           missing_skills=['Leadership', 'Rust', 'GraphQL'],
                           ats_score=ats_score, insight=insight,
                           confidence=min(98, ats_score + 5), decision=decision, job_role=job_role)

@app.route("/pipeline-result")
def pipeline_result():
    if "user" not in session:
        return redirect("/login")
    return render_template("resume_pipeline_layer/results.html", results=session.get("screening_results", []))

@app.route("/results")
def results():
    if "user" not in session:
        return redirect("/login")
    conn = get_screening_db()
    try:
        rows = [dict(r) for r in conn.execute("SELECT * FROM candidates ORDER BY created_at DESC").fetchall()]
        for r in rows:
            try:
                r['skills_json'] = json.loads(r['skills_json']) if r['skills_json'] else []
            except:
                r['skills_json'] = []
    except Exception as e:
        print("Results error:", e)
        rows = []
    conn.close()
    return render_template("results/results.html", results=rows)

@app.route("/ai-feedback")
def ai_feedback_page():
    if "user" not in session:
        return redirect("/login")
    return render_template("ai_modules/ai-feedback.html")

@app.route("/resume-intelligence")
def resume_intelligence():
    if "user" not in session:
        return redirect("/login")
    return render_template("ai_modules/resume-intelligence.html")

@app.route("/ai-screening")
def ai_screening():
    if "user" not in session:
        return redirect("/login")
    return render_template("ai_modules/ai_screening.html")

@app.route("/ai-interviewer")
def ai_interviewer():
    if "user" not in session:
        return redirect("/login")
    conn = get_screening_db()
    try:
        rows = [dict(r) for r in conn.execute("SELECT * FROM candidates ORDER BY score DESC").fetchall()]
        for c in rows:
            try:
                c['skills_json'] = json.loads(c['skills_json']) if c['skills_json'] else []
            except:
                c['skills_json'] = []
    except Exception as e:
        print("Interviewer error:", e)
        rows = []
    conn.close()
    return render_template("ai_modules/ai-interviewer.html", candidates=rows)

@app.route("/compare-candidates", methods=["GET", "POST"])
def compare_candidates_route():
    if "user" not in session:
        return redirect("/login")
    if request.method == "POST":
        data = request.get_json()
        result = compare_candidates(data.get("candidates", []))
        return jsonify(result)
    conn = get_screening_db()
    try:
        rows = [dict(r) for r in conn.execute("SELECT * FROM candidates ORDER BY score DESC").fetchall()]
        for c in rows:
            try:
                c['skills_json'] = json.loads(c['skills_json']) if c['skills_json'] else []
            except:
                c['skills_json'] = []
    except Exception as e:
        print("Compare error:", e)
        rows = []
    conn.close()
    return render_template("ai_modules/compare-candidates.html", candidates=rows)

@app.route("/comparative-ai")
def comparative_ai():
    if "user" not in session:
        return redirect("/login")
    return render_template("comparison_engine/comparative-ai.html")

@app.route("/market-insights")
def market_insights():
    if "user" not in session:
        return redirect("/login")
    return render_template("comparison_engine/market-insights.html")

@app.route("/ranking-analysis")
def ranking_analysis():
    if "user" not in session:
        return redirect("/login")
    return render_template("comparison_engine/ranking-analysis.html")

@app.route("/candidates")
def candidates():
    if "user" not in session:
        return redirect("/login")
    conn = get_screening_db()
    try:
        rows = [dict(r) for r in conn.execute("SELECT * FROM candidates ORDER BY created_at DESC").fetchall()]
        for c in rows:
            try:
                c['skills_json'] = json.loads(c['skills_json']) if c['skills_json'] else []
            except:
                c['skills_json'] = []
    except Exception as e:
        print("Candidates error:", e)
        rows = []
    conn.close()
    return render_template("talent_system/candidates.html", candidates=rows)

@app.route("/candidate-profile")
def candidate_profile():
    if "user" not in session:
        return redirect("/login")
    candidate_id = request.args.get("id")
    conn = get_screening_db()
    try:
        candidate = conn.execute("SELECT * FROM candidates WHERE id=?", (candidate_id,)).fetchone()
        if candidate:
            candidate = dict(candidate)
            try:
                candidate['skills'] = json.loads(candidate['skills_json']) if candidate['skills_json'] else []
            except:
                candidate['skills'] = []
        else:
            candidate = {}
    except Exception as e:
        print("Profile error:", e)
        candidate = {}
    conn.close()
    return render_template("talent_system/candidate-profile.html", candidate=candidate)

@app.route("/candidate-tags")
def candidate_tags():
    if "user" not in session:
        return redirect("/login")
    return render_template("talent_system/candidate-tags.html")

@app.route("/talent-ranking")
def talent_ranking():
    if "user" not in session:
        return redirect("/login")
    return render_template("talent_system/talent-ranking.html")

@app.route("/api/candidates")
def api_candidates():
    if "user" not in session:
        return jsonify({"error": "Unauthorized"}), 401
    conn = get_screening_db()
    try:
        rows = [dict(r) for r in conn.execute("SELECT * FROM candidates").fetchall()]
    except:
        rows = []
    conn.close()
    return jsonify({"candidates": rows})

@app.route("/download-report")
def download_report():
    if "user" not in session:
        return redirect("/login")
    return render_template("reporting_system/download-report.html")

@app.route("/generate-report", methods=["POST"])
def generate_report():
    if "user" not in session:
        return redirect("/login")
    return redirect("/download-report")

@app.route("/api/claude", methods=["POST"])
def claude_proxy():
    if "user" not in session:
        return jsonify({"error": "Unauthorized"}), 401
    data = request.get_json()
    system_prompt = data.get("system", "You are a helpful assistant.")
    messages = data.get("messages", [])
    # Build full prompt
    full_prompt = system_prompt + "\n\n"
    for msg in messages:
        role = "User" if msg["role"] == "user" else "Assistant"
        full_prompt += f"{role}: {msg['content']}\n\n"
    full_prompt += "Assistant:"
    result = call_gemini(full_prompt)
    if result:
        return jsonify({"text": result})
    return jsonify({"error": "Gemini API call failed"}), 500
@app.route("/analytics")
def analytics():
    if "user" not in session:
        return redirect("/login")
    return render_template("analytics/analytics.html")


@app.route("/api/analytics")
def api_analytics():
    if "user" not in session:
        return jsonify({"error": "unauthorized"}), 401
    conn = get_screening_db()
    try:
        rows = conn.execute("SELECT * FROM candidates ORDER BY score DESC").fetchall()
        candidates = [dict(r) for r in rows]
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    finally:
        conn.close()
    total = len(candidates)
    avg_score = sum(c.get("score", 0) or 0 for c in candidates) / total if total else 0

    status_counts = {}
    for c in candidates:
        s = c.get("status") or "pending"
        status_counts[s] = status_counts.get(s, 0) + 1

    buckets = {"0-20": 0, "21-40": 0, "41-60": 0, "61-80": 0, "81-100": 0}
    for c in candidates:
        score = c.get("score", 0) or 0
        if score <= 20: buckets["0-20"] += 1
        elif score <= 40: buckets["21-40"] += 1
        elif score <= 60: buckets["41-60"] += 1
        elif score <= 80: buckets["61-80"] += 1
        else: buckets["81-100"] += 1

    daily_counts = {}
    for c in candidates:
        created = (c.get("created_at") or "")[:10]
        if created:
            daily_counts[created] = daily_counts.get(created, 0) + 1

    return jsonify({
        "total": total,
        "avg_score": round(avg_score, 1),
        "status_counts": status_counts,
        "score_distribution": buckets,
        "daily_counts": daily_counts,
        "top_candidates": candidates[:10]
    })


@app.route("/history")
def history():
    if "user" not in session:
        return redirect("/login")
    conn = get_screening_db()
    try:
        rows = conn.execute("SELECT * FROM candidates ORDER BY created_at DESC LIMIT 50").fetchall()
        history_list = [dict(r) for r in rows]
        for c in history_list:
            try:
                c["skills_json"] = json.loads(c["skills_json"]) if c["skills_json"] else []
            except Exception:
                c["skills_json"] = []
    except Exception:
        history_list = []
    finally:
        conn.close()
    return render_template("history/history.html", history=history_list)

@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username = request.form["username"]
        password = request.form["password"]
        role = request.form.get("role", "recruiter")
        conn = get_db()
        user = conn.execute(
            "SELECT * FROM users WHERE username=? AND password=?",
            (username, password)
        ).fetchone()
        conn.close()
        if user:
            session["user"] = username
            session["role"] = user["role"] if user["role"] else "recruiter"
            # Redirect based on role
            if session["role"] == "jobseeker":
                return redirect("/jobseeker/dashboard")
            return redirect("/dashboard")
        return render_template("authentication/login.html", error="Invalid username or password!")
    return render_template("authentication/login.html")


@app.route("/signup", methods=["GET", "POST"])
def signup():
    if request.method == "POST":
        fullname = request.form["fullname"]
        username = request.form["username"]
        email = request.form["email"]
        phone = request.form.get("phone", "")
        address = request.form.get("address", "")
        password = request.form["password"]
        confirm_password = request.form["confirm_password"]
        role = request.form.get("role", "recruiter")

        if password != confirm_password:
            return render_template("authentication/signup.html", error="Passwords do not match!")

        conn = get_db()
        existing = conn.execute("SELECT * FROM users WHERE username=?", (username,)).fetchone()
        if existing:
            conn.close()
            return render_template("authentication/signup.html", error="Username already taken!")

        conn.execute(
            "INSERT INTO users (fullname, username, email, phone, address, password, role) VALUES (?,?,?,?,?,?,?)",
            (fullname, username, email, phone, address, password, role)
        )
        conn.commit()
        conn.close()
        return render_template("authentication/login.html", success=f"Account created as {role}! Please login.")
    return render_template("authentication/signup.html")


@app.route("/logout")
def logout():
    session.clear()
    return redirect("/")


# ============================================================
# JOBSEEKER DASHBOARD
# ============================================================

@app.route("/jobseeker/dashboard")
def jobseeker_dashboard():
    if "user" not in session or session.get("role") != "jobseeker":
        return redirect("/login?role=jobseeker")
    conn = get_screening_db()
    try:
        # Get this jobseeker's submissions
        my_apps = conn.execute(
            "SELECT * FROM candidates WHERE user_id=? ORDER BY created_at DESC",
            (session["user"],)
        ).fetchall()
        my_apps = [dict(r) for r in my_apps]
        for a in my_apps:
            try:
                a["skills_json"] = json.loads(a["skills_json"]) if a["skills_json"] else []
            except Exception:
                a["skills_json"] = []
    except Exception:
        my_apps = []
    finally:
        conn.close()
    return render_template("jobseeker/dashboard.html", applications=my_apps, username=session["user"])

@app.route("/jobseeker/submit-resume", methods=["POST"])
def jobseeker_submit_resume():
    if "user" not in session:
        return jsonify({"error": "Not logged in"}), 401

    file = request.files.get("resume")
    job_role = request.form.get("job_role", "General Role")

    if not file:
        return jsonify({"error": "No file uploaded"}), 400

    filename = secure_filename(file.filename)
    filepath = os.path.join(UPLOAD_FOLDER, filename)
    file.save(filepath)

    raw_text = ""
    try:
        if filename.lower().endswith(".pdf"):
            import pdfplumber
            with pdfplumber.open(filepath) as pdf:
                for page in pdf.pages:
                    raw_text += (page.extract_text() or "")
        elif filename.lower().endswith(".docx"):
            import docx
            doc = docx.Document(filepath)
            raw_text = "\n".join([p.text for p in doc.paragraphs])
    except Exception as e:
        return jsonify({"error": f"Could not read file: {str(e)}"}), 400

    if not raw_text.strip():
        return jsonify({"error": "Could not extract text from resume"}), 400

    import re
    lines = raw_text.split("\n")
    candidate_name = session["user"]
    for line in lines[:10]:
        line = line.strip()
        if line and len(line) < 50 and not any(c in line for c in ["@", "http", "+"]):
            candidate_name = line
            break

    email_match = re.search(r"[\w\.-]+@[\w\.-]+\.\w+", raw_text)
    candidate_email = email_match.group(0) if email_match else "N/A"

    phone_match = re.search(r"[\+\(]?[0-9][0-9\s\-\(\)]{7,}[0-9]", raw_text)
    candidate_phone = phone_match.group(0).strip() if phone_match else "N/A"

    skill_keywords = [
        "python", "java", "javascript", "react", "node", "sql", "mongodb",
        "html", "css", "git", "docker", "machine learning", "deep learning",
        "tensorflow", "pytorch", "flask", "django", "aws", "azure", "c++",
        "typescript", "vue", "angular", "mysql", "postgresql", "redis",
        "kubernetes", "linux", "excel", "tableau", "power bi", "r", "scala"
    ]
    text_lower = raw_text.lower()
    skills = [skill for skill in skill_keywords if skill in text_lower]

    ats_score = min(100, len(skills) * 8 + (20 if candidate_email != "N/A" else 0) + (10 if len(raw_text) > 500 else 0))

    conn = get_screening_db()
    try:
        conn.execute(
            """INSERT INTO candidates
            (name, email, phone, skills_json, score, status, notes, created_at, updated_at, user_id)
            VALUES (?,?,?,?,?,?,?,?,?,?)""",
            (candidate_name, candidate_email, candidate_phone,
             json.dumps(skills), ats_score, "pending",
             f"Auto-screened for {job_role}",
             datetime.now().isoformat(), datetime.now().isoformat(),
             session["user"])
        )
        conn.commit()
    except Exception as e:
        print(f"DEBUG: DB error: {e}")
    finally:
        conn.close()

    return jsonify({
        "success": True,
        "name": candidate_name,
        "email": candidate_email,
        "score": ats_score,
        "skills": skills,
        "job_role": job_role
    })

# ============================================================
# JOB POSTINGS — RECRUITER
# ============================================================

@app.route("/recruiter/post-job")
def post_job_page():
    if "user" not in session or session.get("role") == "jobseeker":
        return redirect("/login")
    return render_template("recruiter/post-job.html")


@app.route("/api/jobs/post", methods=["POST"])
def api_post_job():
    if "user" not in session:
        return jsonify({"error": "unauthorized"}), 401
    data = request.get_json()
    title = data.get("title", "").strip()
    if not title:
        return jsonify({"error": "Title is required"}), 400

    conn = get_screening_db()
    try:
        conn.execute("""CREATE TABLE IF NOT EXISTS jobs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            title TEXT NOT NULL,
            company TEXT,
            location TEXT,
            job_type TEXT,
            salary TEXT,
            experience TEXT,
            description TEXT,
            skills TEXT,
            min_score INTEGER DEFAULT 0,
            created_by TEXT,
            created_at TEXT,
            is_active INTEGER DEFAULT 1
        )""")
        conn.execute("""
            INSERT INTO jobs (title, company, location, job_type, salary, experience,
                              description, skills, min_score, created_by, created_at, is_active)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,1)
        """, (
            title,
            data.get("company", ""),
            data.get("location", ""),
            data.get("job_type", "Full-time"),
            data.get("salary", ""),
            data.get("experience", ""),
            data.get("description", ""),
            json.dumps(data.get("skills", [])),
            data.get("min_score", 0),
            session["user"],
            datetime.now().isoformat()
        ))
        conn.commit()
        return jsonify({"success": True})
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    finally:
        conn.close()


@app.route("/api/jobs/mine")
def api_my_jobs():
    if "user" not in session:
        return jsonify([])
    conn = get_screening_db()
    try:
        conn.execute("""CREATE TABLE IF NOT EXISTS jobs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            title TEXT, company TEXT, location TEXT, job_type TEXT,
            salary TEXT, experience TEXT, description TEXT, skills TEXT,
            min_score INTEGER DEFAULT 0, created_by TEXT, created_at TEXT, is_active INTEGER DEFAULT 1
        )""")
        # ensure applications table
        conn.execute("""CREATE TABLE IF NOT EXISTS applications (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            job_id INTEGER, applicant TEXT, cover_note TEXT,
            portfolio_url TEXT, status TEXT DEFAULT 'pending',
            created_at TEXT
        )""")
        rows = conn.execute(
            "SELECT * FROM jobs WHERE created_by=? ORDER BY created_at DESC",
            (session["user"],)
        ).fetchall()
        jobs = []
        for r in rows:
            j = dict(r)
            j["skills_list"] = json.loads(j["skills"]) if j.get("skills") else []
            count = conn.execute("SELECT COUNT(*) FROM applications WHERE job_id=?", (j["id"],)).fetchone()[0]
            j["applicant_count"] = count
            jobs.append(j)
        return jsonify(jobs)
    except Exception as e:
        return jsonify([])
    finally:
        conn.close()


@app.route("/api/jobs/<int:job_id>/toggle", methods=["POST"])
def api_toggle_job(job_id):
    if "user" not in session:
        return jsonify({"error": "unauthorized"}), 401
    conn = get_screening_db()
    try:
        job = conn.execute("SELECT is_active FROM jobs WHERE id=? AND created_by=?",
                           (job_id, session["user"])).fetchone()
        if not job:
            return jsonify({"error": "Not found"}), 404
        new_status = 0 if job["is_active"] else 1
        conn.execute("UPDATE jobs SET is_active=? WHERE id=?", (new_status, job_id))
        conn.commit()
        return jsonify({"success": True, "is_active": new_status})
    finally:
        conn.close()


@app.route("/api/jobs/<int:job_id>/delete", methods=["POST"])
def api_delete_job(job_id):
    if "user" not in session:
        return jsonify({"error": "unauthorized"}), 401
    conn = get_screening_db()
    try:
        conn.execute("DELETE FROM jobs WHERE id=? AND created_by=?", (job_id, session["user"]))
        conn.commit()
        return jsonify({"success": True})
    finally:
        conn.close()


# ============================================================
# JOB BROWSING + APPLYING — JOBSEEKER
# ============================================================

@app.route("/jobseeker/jobs")
def jobseeker_jobs_page():
    if "user" not in session or session.get("role") != "jobseeker":
        return redirect("/login")
    return render_template("jobseeker/jobs.html")


@app.route("/api/jobs")
def api_all_jobs():
    conn = get_screening_db()
    try:
        conn.execute("""CREATE TABLE IF NOT EXISTS jobs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            title TEXT, company TEXT, location TEXT, job_type TEXT,
            salary TEXT, experience TEXT, description TEXT, skills TEXT,
            min_score INTEGER DEFAULT 0, created_by TEXT, created_at TEXT, is_active INTEGER DEFAULT 1
        )""")
        conn.execute("""CREATE TABLE IF NOT EXISTS applications (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            job_id INTEGER, applicant TEXT, cover_note TEXT,
            portfolio_url TEXT, status TEXT DEFAULT 'pending', created_at TEXT
        )""")
        rows = conn.execute(
            "SELECT * FROM jobs WHERE is_active=1 ORDER BY created_at DESC"
        ).fetchall()
        jobs = []
        for r in rows:
            j = dict(r)
            j["skills_list"] = json.loads(j["skills"]) if j.get("skills") else []
            count = conn.execute("SELECT COUNT(*) FROM applications WHERE job_id=?", (j["id"],)).fetchone()[0]
            j["applicant_count"] = count
            jobs.append(j)
        return jsonify(jobs)
    except Exception as e:
        return jsonify([])
    finally:
        conn.close()


@app.route("/api/jobs/<int:job_id>/apply", methods=["POST"])
def api_apply_job(job_id):
    if "user" not in session:
        return jsonify({"error": "Not logged in"}), 401
    data = request.get_json() or {}
    conn = get_screening_db()
    try:
        # Check already applied
        existing = conn.execute(
            "SELECT id FROM applications WHERE job_id=? AND applicant=?",
            (job_id, session["user"])
        ).fetchone()
        if existing:
            return jsonify({"error": "Already applied"}), 400
        conn.execute(
            "INSERT INTO applications (job_id, applicant, cover_note, portfolio_url, status, created_at) VALUES (?,?,?,?,?,?)",
            (job_id, session["user"], data.get("cover_note",""), data.get("portfolio_url",""), "pending", datetime.now().isoformat())
        )
        conn.commit()
        return jsonify({"success": True})
    except Exception as e:
        return jsonify({"error": str(e)}), 500
    finally:
        conn.close()



# ============================================================
# RECRUITER — VIEW APPLICATIONS
# ============================================================

@app.route("/recruiter/applications")
def recruiter_applications_page():
    if "user" not in session or session.get("role") == "jobseeker":
        return redirect("/login")
    conn = get_screening_db()
    try:
        rows = conn.execute("""
            SELECT a.*, j.title as job_title, j.company
            FROM applications a
            JOIN jobs j ON a.job_id = j.id
            WHERE j.created_by=?
            ORDER BY a.created_at DESC
        """, (session["user"],)).fetchall()
        apps = [dict(r) for r in rows]
    except Exception:
        apps = []
    finally:
        conn.close()
    return jsonify(apps)  # returns JSON for now; swap for render_template later

# ─────────────────────────────────────────────
# PASTE THESE ROUTES INTO app.py
# Place BEFORE the  if __name__ == "__main__":  line
# ─────────────────────────────────────────────

# Make sure this import is at the top of app.py:
#   from datetime import datetime

# ── PAGE ROUTES ───────────────────────────────


@app.route("/jobseeker/my-applications")
def jobseeker_my_applications_page():
    if "user" not in session or session.get("role") == "recruiter":
        return redirect("/login")
    return render_template("jobseeker/my-applications.html")


# ── API: Recruiter sees all applications for their jobs ───

@app.route("/api/recruiter/applications")
def api_recruiter_applications():
    if "user" not in session:
        return jsonify([])
    try:
        conn = sqlite3.connect("users.db")
        conn.row_factory = sqlite3.Row
        cur = conn.cursor()

        # Get all applications for jobs posted by this recruiter
        cur.execute("""
            SELECT
                a.id,
                a.job_id,
                a.applicant,
                a.ats_score,
                a.cover_note,
                a.portfolio_url,
                a.status,
                a.created_at,
                j.title   AS job_title,
                j.company AS company,
                j.location AS location
            FROM applications a
            JOIN jobs j ON a.job_id = j.id
            WHERE j.created_by = ?
            ORDER BY a.created_at DESC
        """, (session["user"],))

        rows = [dict(r) for r in cur.fetchall()]
        conn.close()
        return jsonify(rows)
    except Exception as e:
        return jsonify({"error": str(e)}), 500


# ── API: Jobseeker sees their own applications ────────────

@app.route("/api/my-applications")
def api_my_applications():
    if "user" not in session:
        return jsonify([])
    try:
        conn = sqlite3.connect("users.db")
        conn.row_factory = sqlite3.Row
        cur = conn.cursor()

        cur.execute("""
            SELECT
                a.id,
                a.job_id,
                a.ats_score,
                a.cover_note,
                a.portfolio_url,
                a.status,
                a.created_at,
                j.title    AS job_title,
                j.company  AS company,
                j.location AS location
            FROM applications a
            JOIN jobs j ON a.job_id = j.id
            WHERE a.applicant = ?
            ORDER BY a.created_at DESC
        """, (session["user"],))

        rows = [dict(r) for r in cur.fetchall()]
        conn.close()
        return jsonify(rows)
    except Exception as e:
        return jsonify({"error": str(e)}), 500


# ── API: Recruiter updates application status ─────────────

@app.route("/api/applications/<int:app_id>/status", methods=["POST"])
def api_update_application_status(app_id):
    if "user" not in session or session.get("role") == "jobseeker":
        return jsonify({"error": "Unauthorized"}), 403
    data = request.get_json() or {}
    status = data.get("status", "pending")
    if status not in ("pending", "shortlisted", "rejected"):
        return jsonify({"error": "Invalid status"}), 400
    try:
        conn = sqlite3.connect("users.db")
        conn.execute(
            "UPDATE applications SET status=? WHERE id=?",
            (status, app_id)
        )
        conn.commit()
        conn.close()
        return jsonify({"success": True})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


# ── DB SETUP: make sure applications table has status column ──
# Run this once — add to your DB init section or run manually:
#
#   conn = sqlite3.connect("users.db")
#   conn.execute("""
#       CREATE TABLE IF NOT EXISTS applications (
#           id INTEGER PRIMARY KEY AUTOINCREMENT,
#           job_id INTEGER NOT NULL,
#           applicant TEXT NOT NULL,
#           ats_score INTEGER DEFAULT 0,
#           cover_note TEXT,
#           portfolio_url TEXT,
#           status TEXT DEFAULT 'pending',
#           created_at TEXT DEFAULT (datetime('now','localtime')),
#           FOREIGN KEY(job_id) REFERENCES jobs(id)
#       )
#   """)
#   try:
#       conn.execute("ALTER TABLE applications ADD COLUMN status TEXT DEFAULT 'pending'")
#   except:
#       pass  # column already exists
#   conn.commit()
#   conn.close()

if __name__ == "__main__":
    app.run(debug=True)
